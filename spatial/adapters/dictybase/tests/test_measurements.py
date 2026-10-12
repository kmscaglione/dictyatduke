"""The authors' fractionation matrix and figure: exact storage, identity with
Table S2, the reconstructed map, and reproducible computed maps.

Synthetic inputs cover the code paths everywhere. LocalDataTest checks the real
files and skips where they are absent (they are never committed)."""
import collections
import json
import os
import pathlib
import subprocess
import sys
import tempfile
import unittest
import zlib

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE.parents[2] / "py"))
import build_bundle as bb  # noqa: E402
import figure_coordinates as fc  # noqa: E402
from spatialprot import bundle as B, summary, validate as V  # noqa: E402

try:
    import numpy as np
    import sklearn  # noqa: F401
    import umap  # noqa: F401
    from spatialprot import embed
    LIBS = True
except ImportError:  # pragma: no cover
    LIBS = False

LOCAL = bb.DEFAULT_OUT.parent
SOURCE = LOCAL / "source"
MATRIX = SOURCE / "DDPCP_raw_all_nbavg_norm_zero_cut51.csv"
FIGURE = SOURCE / "Tsne plot with coordinates.pdf"
XLSX = SOURCE / "media-1.xlsx"
REPORT = LOCAL / "validation_report.json"


def fake_figure(points, overlay, x_ticks=(-40, -20, 0, 20, 40), y_ticks=(-40, -20, 0, 20, 40)):
    """A minimal R-style vector figure. points: [(x, y)] in data units, drawn grey
    in order; overlay: [(index into points, (r, g, b))] drawn again in colour."""
    ax, bx, ay, by = 4.5, 286.0, 5.2, 283.0                       # data -> PDF units
    px = lambda v: ax * v + bx
    py = lambda v: ay * v + by
    out = []
    for t in x_ticks:
        out.append("%.2f 72.00 m %.2f 64.80 l  S" % (px(t), px(t)))
    for t in x_ticks:
        out += ["BT", "/F2 1 Tf 12.00 0.00 0.00 12.00 %.2f 46.08 Tm (%d) Tj" % (px(t) - 8, t), "ET"]
    for t in y_ticks:
        out.append("57.60 %.2f m 50.40 %.2f l  S" % (py(t), py(t)))
    for t in y_ticks:
        out += ["BT", "/F2 1 Tf 0.00 12.00 -12.00 0.00 43.20 %.2f Tm (%d) Tj" % (py(t) - 8, t), "ET"]
    out += ["BT", "/F2 1 Tf 12.00 0.00 0.00 12.00 247.46 17.28 Tm (Dimension 1) Tj", "ET"]

    def circle(x, y):
        cx, cy, r = px(x), py(y), 3.24
        return ["  %.2f %.2f m" % (cx - r, cy), "  %.2f %.2f %.2f %.2f %.2f %.2f c" % (cx - r, cy + 1.79, cx - 1.79, cy + r, cx, cy + r),
                "  %.2f %.2f %.2f %.2f %.2f %.2f c" % (cx + 1.79, cy + r, cx + r, cy + 1.79, cx + r, cy),
                "  %.2f %.2f %.2f %.2f %.2f %.2f c" % (cx + r, cy - 1.79, cx + 1.79, cy - r, cx, cy - r),
                "  %.2f %.2f %.2f %.2f %.2f %.2f c" % (cx - 1.79, cy - r, cx - r, cy - 1.79, cx - r, cy), "B"]
    out += ["/sRGB cs", "0.859 0.855 0.855 scn"]
    for x, y in points:
        out += circle(x, y)
    last = None
    for i, rgb in overlay:
        if rgb != last:
            out.append("%.3f %.3f %.3f scn" % rgb)
            last = rgb
        out += circle(*points[i])
    body = zlib.compress(("q 57.60 72.00 446.40 432.00 re W n\n" + "\n".join(out) + "\nQ\n").encode("latin-1"))
    return b"%PDF-1.4\n1 0 obj\n<< /Length " + str(len(body)).encode() + b" /Filter /FlateDecode >>\nstream\n" + body + b"\nendstream\nendobj\n"


class FigureCoordinatesTest(unittest.TestCase):
    PTS = [(-30.5, 12.25), (0.0, 0.0), (17.75, -22.5), (39.0, 38.0), (17.75, -22.5), (-5.0, 5.0)]
    IDS = ["A", "B", "C", "D", "E", "F"]
    RED, BLUE = (0.894, 0.102, 0.110), (0.216, 0.494, 0.722)

    def test_positions_and_row_identity_are_recovered(self):
        pdf = fake_figure(self.PTS, [(0, self.RED), (3, self.RED), (5, self.BLUE)])
        coords, groups, facts = fc.recover(pdf, self.IDS)
        self.assertEqual(list(coords), self.IDS)
        for rid, (x, y) in zip(self.IDS, self.PTS):
            self.assertAlmostEqual(coords[rid][0], x, delta=2 * facts["position_resolution"] + 0.002)
            self.assertAlmostEqual(coords[rid][1], y, delta=0.01)
        self.assertEqual(dict(groups), {"#e41a1c": ["A", "D"], "#377eb8": ["F"]})
        self.assertEqual((facts["base_points"], facts["overlay_points"], facts["overlay_points_unmatched"]), (6, 3, 0))
        self.assertLess(facts["axis_calibration_worst_residual"], 0.01)

    def test_shared_positions_are_reported_not_guessed(self):
        pdf = fake_figure(self.PTS, [(2, self.RED)])                 # C and E sit on the same spot
        coords, groups, facts = fc.recover(pdf, self.IDS)
        self.assertEqual(facts["positions_shared_by_several_rows"], [["C", "E"]])
        self.assertEqual((facts["overlay_points_ambiguous"], facts["distinct_positions"]), (1, 5))
        self.assertEqual(dict(groups), {})                           # the ambiguous coloured point is assigned to neither
        self.assertEqual(coords["C"], coords["E"])

    def test_wrong_row_count_or_unreadable_figure_stops(self):
        pdf = fake_figure(self.PTS, [])
        with self.assertRaises(ValueError):
            fc.recover(pdf, self.IDS + ["G"])
        with self.assertRaises(ValueError):
            fc.recover(b"%PDF-1.4 nothing here", self.IDS)
        with self.assertRaises(ValueError):
            fc.calibrate({"x": [1.0], "y": [1.0, 2.0], "labels": ["0", "0", "1"]})


class MatrixReadTest(unittest.TestCase):
    HEADER = ["Accession"] + [("" if f == "SUP" else "X") + f"{f}.{k}" for k in (1, 2, 3) for f in bb.FRACTION_LABELS]

    def write(self, rows, header=None):
        import csv
        fh = tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, newline="")
        csv.writer(fh).writerows([header or self.HEADER] + rows)
        fh.close()
        self.addCleanup(os.unlink, fh.name)
        return fh.name

    def test_values_are_kept_exactly(self):
        vals = ["0.524541623", "1.34E-08", "0"] + ["0.1"] * 27
        columns, ids, values, text = bb.read_matrix(self.write([["P1"] + vals]))
        self.assertEqual((len(columns), ids), (30, ["P1"]))
        self.assertEqual(text["P1"], vals)
        self.assertEqual(values["P1"][:3], [0.524541623, 1.34e-08, 0.0])
        self.assertEqual([repr(v) for v in values["P1"]], [repr(float(t)) for t in vals])

    def test_bad_files_are_refused(self):
        row = ["P1"] + ["0.1"] * 30
        with self.assertRaises(SystemExit):
            bb.read_matrix(self.write([row], header=self.HEADER[:-1] + ["SUP.4"]))
        with self.assertRaises(SystemExit):
            bb.read_matrix(self.write([row, row]))                   # the same protein group twice
        with self.assertRaises(SystemExit):
            bb.read_matrix(self.write([["P1"] + ["0.1"] * 29 + [""]]))

    def test_fractions_keep_three_conditions_of_ten(self):
        fr = bb.fraction_list()
        self.assertEqual(len(fr), 30)
        self.assertEqual(collections.Counter(f["condition"] for f in fr), {"Lysis condition 1": 10, "Lysis condition 2": 10, "Lysis condition 3": 10})
        self.assertEqual([f["label"] for f in fr[:10]], ["1K", "3K", "5K", "9K", "12K", "15K", "30K", "79K", "120K", "SUP"])
        self.assertEqual(len({f["id"] for f in fr}), 30)


@unittest.skipUnless(LIBS, "NumPy, scikit-learn and umap-learn are needed")
class AddMeasurementsTest(unittest.TestCase):
    """The whole step on a small invented dataset."""

    @classmethod
    def setUpClass(cls):
        import csv
        import random
        rng = random.Random(5)
        classes = ["Mitochondria", "Cytosol", "Nucleus"]
        peak = {"Mitochondria": 0, "Cytosol": 9, "Nucleus": 1}
        svm, rows = [], []
        for i in range(60):
            acc, c = "XP_%04d.1" % i, classes[i % 3]
            prof = []
            for _ in range(3):
                v = [rng.random() * 0.05 for _ in range(10)]
                v[peak[c]] += 1.0
                t = sum(v)
                v = [round(x / t, 9) for x in v]
                v[peak[c]] = round(1 - (sum(v) - v[peak[c]]), 9)
                prof += v
            rows.append([acc] + [repr(x) for x in prof])
            score = round(0.3 + 0.01 * i, 4)
            svm.append((acc, c, score, c if i % 2 else "unknown"))
        cls.tables = {"svm": svm, "markers": [(svm[1][0], svm[1][1], "Yes"), (svm[3][0], svm[3][1], "No")], "mito": [(svm[3][0], "Overlapping Proteins")]}
        fh = tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, newline="")
        csv.writer(fh).writerows([MatrixReadTest.HEADER] + sorted(rows))
        fh.close()
        cls.matrix = fh.name
        cls.rows = {r[0]: r[1:] for r in rows}
        p2g = {r[0]: ("DDB_G%07d" % i, None) for i, r in enumerate(svm)}
        cls.plain = bb.build(cls.tables, p2g, {}, [{"name": "t", "description": "t"}], built="2026-01-01")
        cls.b = bb.build(cls.tables, p2g, {}, [{"name": "t", "description": "t"}], built="2026-01-01")
        cls.report = bb.add_measurements(cls.b, cls.tables, cls.matrix, None, "2026-10-12")

    @classmethod
    def tearDownClass(cls):
        os.unlink(cls.matrix)

    def test_bundle_is_valid_and_published_layers_are_untouched(self):
        self.assertEqual(V.errors(V.validate(self.b)), [])
        self.assertEqual(self.b["layers"], self.plain["layers"])
        self.assertTrue(self.report["published_layers_unchanged"])
        self.assertEqual(self.b["entities"], self.plain["entities"])

    def test_every_value_is_the_supplied_value(self):
        vals = self.b["profiles"]["values"]
        self.assertEqual(set(vals), set(self.rows))
        for acc, text in self.rows.items():
            self.assertEqual(vals[acc], [float(t) for t in text])
        again = json.loads(B.dumps(self.b))["profiles"]["values"]
        self.assertEqual(again, vals)                                 # and it survives being written
        self.assertTrue(self.report["matrix"]["values_parsed_equal_supplied_text"])
        self.assertIn("No further normalization or imputation", self.b["profiles"]["processing"])
        self.assertEqual(self.b["profiles"]["evidence_type"], "measured_profile")

    def test_three_computed_maps_cover_exactly_the_matrix_rows(self):
        self.assertEqual([e["id"] for e in self.b["embeddings"]], ["tsne", "pca", "umap"])
        for e in self.b["embeddings"]:
            self.assertEqual(set(e["coordinates"]), set(self.rows), e["id"])
            self.assertEqual(e["source"], "computed")
            self.assertIn("random_state" if e["id"] != "pca" else "explained_variance", e["method"]["parameters"])

    def test_rows_that_are_not_table_s2_stop_the_build(self):
        other = dict(self.tables, svm=self.tables["svm"][:-1])
        with self.assertRaises(SystemExit):
            bb.add_measurements(bb.build(other, {}, {}, [{"name": "t", "description": "t"}]), other, self.matrix, None, "2026-10-12")

    def test_class_means_are_reported(self):
        peaks = self.report["class_mean_profiles"]
        self.assertEqual(peaks["Cytosol"]["peak_fraction_per_condition"], ["SUP", "SUP", "SUP"])
        self.assertEqual(peaks["Mitochondria"]["peak_fraction_per_condition"], ["1K", "1K", "1K"])


@unittest.skipUnless(MATRIX.exists() and REPORT.exists() and bb.DEFAULT_OUT.exists(), "local dataset not built (it is never committed)")
class LocalDataTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.b = B.load(bb.DEFAULT_OUT)
        cls.report = json.loads(REPORT.read_text())
        cls.columns, cls.ids, cls.values, cls.text = bb.read_matrix(MATRIX)

    def test_matrix_is_stored_complete_and_unchanged(self):
        vals = self.b["profiles"]["values"]
        self.assertEqual(len(vals), 6337)
        self.assertEqual(len(self.b["fractions"]), 30)
        for acc in self.ids:
            self.assertEqual(vals[acc], [float(t) for t in self.text[acc]], acc)
        self.assertEqual(sum(len(v) for v in vals.values()), 6337 * 30)

    def test_rows_are_exactly_the_detected_groups_of_table_s2(self):
        detected = {e["id"] for e in self.b["entities"] if e.get("detected", True)}
        self.assertEqual(set(self.ids), detected)
        self.assertEqual(set(self.ids), {a["entity"] for a in summary.layer(self.b, "svm")["assignments"]})
        undetected = {e["id"] for e in self.b["entities"]} - detected
        self.assertEqual(len(undetected), 14)
        self.assertFalse(undetected & set(self.b["profiles"]["values"]))      # nothing invented for them

    def test_each_lysis_condition_sums_to_one(self):
        for acc in self.ids:
            v = self.values[acc]
            for k in range(3):
                self.assertAlmostEqual(sum(v[k * 10:(k + 1) * 10]), 1.0, places=5, msg=acc)

    def test_sources_are_recorded_with_checksums(self):
        src = {s["name"]: s for s in self.b["dataset"]["provenance"]["sources"]}
        for path in (MATRIX, FIGURE, XLSX):
            self.assertEqual(src[path.name]["sha256"], B.sha256_file(path), path.name)
        self.assertEqual(self.report["received"], "2026-10-12")

    def test_published_results_are_untouched(self):
        import openpyxl
        wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
        rows = [r for r in list(wb["Table S2 SVM Scores"].iter_rows(values_only=True))[1:] if r and r[0]]
        svm = {a["entity"]: a for a in summary.layer(self.b, "svm")["assignments"]}
        label = {c["id"]: c["label"] for c in self.b["compartments"]}
        for acc, klass, score, pred in rows:
            a = svm[acc]
            self.assertEqual((label[a["compartment"]], a["score"], a["status"] == "assigned"), (klass, score, pred != "unknown"), acc)
        self.assertNotIn("threshold", summary.layer(self.b, "svm")["score"])
        self.assertEqual([l["id"] for l in self.b["layers"]], ["svm", "markers-training", "markers-heldout", "mito-compendium"])

    def test_every_map_point_is_a_matrix_row(self):
        self.assertEqual([e["id"] for e in self.b["embeddings"]], ["tsne", "pca", "umap", "authors-tsne"])
        for e in self.b["embeddings"]:
            self.assertEqual(set(e["coordinates"]), set(self.ids), e["id"])
        self.assertEqual([e["source"] for e in self.b["embeddings"]], ["computed", "computed", "computed", "author"])

    @unittest.skipUnless(LIBS, "NumPy is needed")
    def test_pca_is_reproduced_exactly(self):
        stored = next(e for e in self.b["embeddings"] if e["id"] == "pca")
        self.assertEqual(embed.pca(self.b), stored)
        ev = stored["method"]["parameters"]["explained_variance"]
        self.assertEqual(self.report["computed_maps"]["pca"]["explained_variance"], ev)

    @unittest.skipUnless(LIBS and os.environ.get("SPATIAL_SLOW") == "1", "set SPATIAL_SLOW=1 to recompute t-SNE and UMAP (about a minute)")
    def test_tsne_and_umap_are_reproduced_exactly(self):
        for fn, name in ((embed.tsne, "tsne"), (embed.umap, "umap")):
            stored = next(e for e in self.b["embeddings"] if e["id"] == name)
            self.assertEqual(fn(self.b)["coordinates"], stored["coordinates"], name)

    def test_reconstructed_map_is_labelled_and_passes_its_checks(self):
        rec = next(e for e in self.b["embeddings"] if e["id"] == "authors-tsne")
        self.assertEqual(rec["label"], "Authors' t-SNE — coordinates reconstructed from published figure")
        self.assertIn("not their original numbers", rec["note"])
        self.assertIn("not used in any calculation", rec["note"])
        f = self.report["authors_map_reconstruction"]
        self.assertEqual((f["base_points"], f["overlay_points_unmatched"]), (6337, 0))
        self.assertGreaterEqual(f["coloured_points_whose_group_matches_column_svm"], 0.99)
        self.assertLess(f["same_test_under_table_s2_row_order"], 0.3)
        self.assertGreater(f["mean_profile_correlation_with_nearest_map_neighbour"] - f["same_with_rows_shuffled"], 0.5)
        self.assertGreaterEqual(f["trustworthiness_k15_against_supplied_matrix"], 0.9)
        self.assertTrue(f["judged_reliable_for_viewing"])
        coords, groups, facts = fc.recover(FIGURE.read_bytes(), self.ids)
        self.assertEqual(coords, rec["coordinates"])                  # recovery itself is deterministic
        self.assertEqual(facts["distinct_positions"], f["distinct_positions"])

    def test_figure_and_table_differences_are_recorded_not_resolved(self):
        f = self.report["figure_versus_table_s2"]
        self.assertEqual(f["final_calls_in_table"], 3169)
        self.assertNotEqual(f["coloured_points"], f["final_calls_in_table"])
        self.assertEqual(sum(v["final_calls_svm_pred"] for v in f["classes"].values()), 3169)

    def test_restricted_files_are_not_tracked(self):
        out = subprocess.run(["git", "ls-files", "--", str(LOCAL)], cwd=bb.REPO, capture_output=True, text=True).stdout
        self.assertEqual(out.strip(), "")
        for path in (MATRIX, FIGURE, REPORT, bb.DEFAULT_OUT):
            self.assertEqual(subprocess.run(["git", "check-ignore", "-q", str(path)], cwd=bb.REPO).returncode, 0, path.name)


if __name__ == "__main__":
    unittest.main()
